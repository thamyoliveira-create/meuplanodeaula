#!/usr/bin/env python3
"""Merge the existing 2nd-series EPT scope with a 3rd-series workbook.

The browser only receives fields used by the lesson-plan generator. This keeps
the generated JavaScript compact even when the source workbook repeats long
competency descriptions on every lesson row.
"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string


SHEET_COURSES = {
    "ADM": "Administração",
    "AGRO": "Agronegócio",
    "DADOS": "Ciência de Dados",
    "ENF": "Enfermagem",
    "FARM": "Farmácia",
    "HOTEL": "Hospedagem",
    "LOG": "Logística",
    "PMD": "Comum",
    "SIS": "Desenvolvimento de Sistemas",
    "VEND": "Vendas",
}

OUTPUT_FIELDS = (
    "_year",
    "_course",
    "Bimestre",
    "Componente de 3 ou 4 aulas semanais?",
    "Nome do componente",
    "Semana",
    "Tema da semana",
    "Título da aula",
    "Habilidades técnicas",
    "Habilidades socioemocionais",
    "Objetivos da aula",
    "Objeto de conhecimento",
)


def clean(value: object) -> str:
    if value is None:
        return ""
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text


def normalized_header(value: object) -> str:
    text = clean(value).lower()
    text = "".join(
        character for character in unicodedata.normalize("NFD", text)
        if unicodedata.category(character) != "Mn"
    )
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def first_value(row: dict, names: tuple[str, ...]) -> str:
    for name in names:
        value = clean(row.get(name, ""))
        if value:
            return value
    return ""


def compact_existing_row(row: dict) -> dict:
    return {
        "_year": clean(row.get("_year")),
        "_course": clean(row.get("_course")),
        "Bimestre": first_value(row, ("Bimestre",)),
        "Componente de 3 ou 4 aulas semanais?": first_value(
            row, ("Componente de 3 ou 4 aulas semanais?",)
        ),
        "Nome do componente": first_value(row, ("Nome do componente",)),
        "Semana": first_value(row, ("Semana",)),
        "Tema da semana": first_value(row, ("Tema da semana",)),
        "Título da aula": first_value(row, ("Título da aula",)),
        "Habilidades técnicas": first_value(
            row, ("Habilidades técnicas", "Habilidade técnica")
        ),
        "Habilidades socioemocionais": first_value(
            row,
            (
                "Habilidades socioemocionais",
                "Habildades socioemocionais",
                "Competências Socioemocionais",
                "Competências socioemocionais",
            ),
        ),
        "Objetivos da aula": first_value(
            row, ("Objetivos da aula", "Objetivo da aula")
        ),
        "Objeto de conhecimento": first_value(
            row,
            (
                "Objeto de conhecimento",
                "Objeto de conhecimento ",
                "Objeto de conhecimento – macro",
                "Objeto de conhecimento - macro",
                "Objeto de conhecimento macro",
            ),
        ),
    }


def load_existing_second_series(path: Path) -> list[dict]:
    source = path.read_text(encoding="utf-8")
    match = re.search(r"window\.CURRICULUM_DATA\s*=\s*(\[.*\])\s*;?\s*$", source, re.DOTALL)
    if not match:
        raise ValueError(f"Could not parse curriculum data from {path}")
    rows = json.loads(match.group(1))
    return [compact_existing_row(row) for row in rows if clean(row.get("_year")) == "2º Ano (MTEC)"]


CELL_REFERENCE_RE = re.compile(
    r"^=(?:(?:'(?P<quoted>[^']+)'|(?P<plain>[^!]+))!)?\$?(?P<column>[A-Z]+)\$?(?P<row>\d+)$"
)


def resolved_cell_value(workbook, worksheet, row: int, column: int, seen=None):
    seen = seen or set()
    key = (worksheet.title, row, column)
    if key in seen:
        return ""
    seen.add(key)

    value = worksheet.cell(row, column).value
    if not isinstance(value, str) or not value.startswith("="):
        return value
    match = CELL_REFERENCE_RE.match(value.strip())
    if not match:
        return ""
    sheet_name = match.group("quoted") or match.group("plain") or worksheet.title
    target = workbook[sheet_name]
    return resolved_cell_value(
        workbook,
        target,
        int(match.group("row")),
        column_index_from_string(match.group("column")),
        seen,
    )


def find_column(headers: dict[str, int], *names: str) -> int | None:
    for name in names:
        column = headers.get(normalized_header(name))
        if column:
            return column
    return None


def extract_third_series(path: Path) -> list[dict]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    rows: list[dict] = []

    for sheet_name, course in SHEET_COURSES.items():
        if sheet_name not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_name]
        headers = {
            normalized_header(worksheet.cell(1, column).value): column
            for column in range(1, worksheet.max_column + 1)
            if normalized_header(worksheet.cell(1, column).value)
        }

        columns = {
            "bimestre": find_column(headers, "Bimestre"),
            "hours": find_column(headers, "Componente de 3 ou 4 aulas semanais?"),
            "component": find_column(headers, "Nome do componente"),
            "week": find_column(headers, "Semana"),
            "theme": find_column(headers, "Tema da semana"),
            "title": find_column(headers, "Título da aula"),
            "technical": find_column(headers, "Habilidades técnicas"),
            "social": find_column(
                headers, "Habilidades socioemocionais", "Habildades socioemocionais"
            ),
            "objective": find_column(headers, "Objetivos da aula", "Objetivo da aula"),
            "knowledge": find_column(
                headers,
                "Objeto de conhecimento",
                "Objeto de conhecimento macro",
                "Objeto de conhecimento - macro",
                "Objeto de conhecimento – macro",
            ),
        }
        missing = [name for name, column in columns.items() if column is None]
        if missing:
            raise ValueError(f"Missing columns in {sheet_name}: {', '.join(missing)}")

        for row_number in range(2, worksheet.max_row + 1):
            values = {
                name: resolved_cell_value(workbook, worksheet, row_number, column)
                for name, column in columns.items()
            }
            component = clean(values["component"])
            title = clean(values["title"])
            try:
                week = int(float(values["week"]))
            except (TypeError, ValueError):
                continue
            if not component or not title or not 1 <= week <= 28:
                continue

            try:
                bimestre = int(float(values["bimestre"]))
            except (TypeError, ValueError):
                bimestre = min(4, ((week - 1) // 7) + 1)
            try:
                weekly_hours = int(float(values["hours"]))
            except (TypeError, ValueError):
                weekly_hours = 3

            rows.append({
                "_year": "3º Ano (MTEC)",
                "_course": course,
                "Bimestre": str(bimestre),
                "Componente de 3 ou 4 aulas semanais?": str(weekly_hours),
                "Nome do componente": component,
                "Semana": str(week),
                "Tema da semana": clean(values["theme"]),
                "Título da aula": title,
                "Habilidades técnicas": clean(values["technical"]),
                "Habilidades socioemocionais": clean(values["social"]),
                "Objetivos da aula": clean(values["objective"]),
                "Objeto de conhecimento": clean(values["knowledge"]),
            })

    return rows


def deduplicate(rows: list[dict]) -> list[dict]:
    unique: dict[tuple[str, ...], dict] = {}
    for row in rows:
        key = (
            row["_year"], row["_course"], row["Nome do componente"],
            row["Semana"], row["Título da aula"],
        )
        unique[key] = {field: clean(row.get(field, "")) for field in OUTPUT_FIELDS}
    return sorted(
        unique.values(),
        key=lambda row: (
            row["_year"], row["_course"], row["Nome do componente"],
            int(row["Semana"]), row["Título da aula"],
        ),
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("existing_js", type=Path)
    parser.add_argument("third_series_xlsx", type=Path)
    parser.add_argument("output_js", type=Path)
    args = parser.parse_args()

    rows = deduplicate(
        load_existing_second_series(args.existing_js)
        + extract_third_series(args.third_series_xlsx)
    )
    args.output_js.write_text(
        "window.CURRICULUM_DATA = "
        + json.dumps(rows, ensure_ascii=False, separators=(",", ":"))
        + ";\n",
        encoding="utf-8",
    )

    courses: dict[str, int] = {}
    for row in rows:
        key = f'{row["_year"]} | {row["_course"]}'
        courses[key] = courses.get(key, 0) + 1
    print(json.dumps({"rows": len(rows), "courses": courses}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
